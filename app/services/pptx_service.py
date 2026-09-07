"""
Sales analysis PowerPoint generation, with real python-pptx chart
objects (not just text tables) for payment mix and top products, plus
a stock-health slide so low-stock items are visible at a glance.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

from app.config import REPORT_DIR
from app.services.sales_service import get_daily_sales_summary
from app.tools.inventory_tools import get_low_stock


def _add_bar_chart(slide, categories, values, series_name, left, top, width, height):
    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series(series_name, values)

    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        left, top, width, height,
        chart_data
    )

    chart = graphic_frame.chart
    chart.has_legend = False

    return chart


def _add_pie_chart(slide, categories, values, series_name, left, top, width, height):
    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series(series_name, values)

    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.PIE,
        left, top, width, height,
        chart_data
    )

    chart = graphic_frame.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False

    return chart


def generate_sales_pptx(output_path=None, target_date=None):
    """
    Generate a PowerPoint sales analysis report for the specified
    date (defaults to today).
    """

    summary = get_daily_sales_summary(target_date)

    if not summary["success"]:
        return {"success": False, "message": summary["message"]}

    if output_path is None:
        output_path = str(REPORT_DIR / f"sales_analysis_{summary['date']}.pptx")

    prs = Presentation()

    # -------------------------------------------------
    # Slide 1 - Title
    # -------------------------------------------------

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Supermarket Sales Analysis"
    slide.placeholders[1].text = f"Daily Sales Report\nDate: {summary['date']}"

    # -------------------------------------------------
    # Slide 2 - Daily Overview
    # -------------------------------------------------

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Daily Sales Overview"

    body = slide.placeholders[1]
    body.text = (
        f"Total Bills: {summary['total_bills']}\n"
        f"Total Sales: Rs. {summary['total_sales']:.2f}\n"
        f"Taxable Amount: Rs. {summary['total_subtotal']:.2f}\n"
        f"CGST: Rs. {summary['total_cgst']:.2f}\n"
        f"SGST: Rs. {summary['total_sgst']:.2f}\n"
        f"Total GST: Rs. {summary['total_gst']:.2f}"
    )

    # -------------------------------------------------
    # Slide 3 - Payment Analysis (real pie chart)
    # -------------------------------------------------

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Payment Analysis"

    payment = summary["payment_breakdown"]
    nonzero_payment = {k: v for k, v in payment.items() if v > 0}

    if nonzero_payment:
        _add_pie_chart(
            slide,
            list(nonzero_payment.keys()),
            list(nonzero_payment.values()),
            "Payment Mix",
            Inches(1.5), Inches(1.4), Inches(7), Inches(5)
        )
    else:
        textbox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
        textbox.text_frame.text = "No finalized sales were recorded for this date."

    # -------------------------------------------------
    # Slide 4 - Top Products (real bar chart)
    # -------------------------------------------------

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Top Selling Products"

    products = summary["top_products"][:8]

    if products:
        _add_bar_chart(
            slide,
            [p["product"] for p in products],
            [p["quantity"] for p in products],
            "Units Sold",
            Inches(0.8), Inches(1.4), Inches(8.4), Inches(5)
        )
    else:
        textbox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
        textbox.text_frame.text = "No finalized sales were recorded for this date."

    # -------------------------------------------------
    # Slide 5 - Stock Health
    # -------------------------------------------------

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Stock Health (Low / Reorder-Level Items)"

    low_stock = get_low_stock()
    low_stock_products = low_stock.get("products", []) if low_stock.get("success") else []

    if low_stock_products:
        rows = min(len(low_stock_products), 10) + 1
        cols = 3

        table = slide.shapes.add_table(
            rows, cols, Inches(1), Inches(1.4), Inches(8), Inches(4.5)
        ).table

        table.cell(0, 0).text = "Product"
        table.cell(0, 1).text = "Current Stock"
        table.cell(0, 2).text = "Reorder Level"

        for i, product in enumerate(low_stock_products[:10], start=1):
            table.cell(i, 0).text = product["product"]
            table.cell(i, 1).text = f"{product['quantity']} {product['unit']}"
            table.cell(i, 2).text = str(product["reorder_level"])
    else:
        textbox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
        textbox.text_frame.text = "All products are above their reorder level."

    # -------------------------------------------------
    # Slide 6 - Daily Closing Summary
    # -------------------------------------------------

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Daily Closing Summary"

    body = slide.placeholders[1]
    body.text = (
        f"Business Date: {summary['date']}\n"
        f"Bills Generated: {summary['total_bills']}\n"
        f"Revenue: Rs. {summary['total_sales']:.2f}\n"
        f"GST Collected: Rs. {summary['total_gst']:.2f}\n\n"
        "Report generated by Supermarket Ops Agent"
    )

    prs.save(output_path)

    return {
        "success": True,
        "message": "Sales analysis PPTX generated successfully.",
        "file": output_path,
        "date": summary["date"]
    }
